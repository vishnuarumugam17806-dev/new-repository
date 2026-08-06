"""
SMART DB — Schema Generator Blueprint
Handles schema creation, database generation, sql downloads, DBMS execution and deployment.
"""
import os
import re
import io
import csv
import json
import logging
from urllib.parse import quote_plus
from flask import Blueprint, render_template, request, redirect, url_for, flash, Response, jsonify, send_file
from flask_login import current_user
import config as cfg
from models import db, SharedDatabase, SchemaVersion, DownloadLog, IndustryTemplate, Comment
from ai_engine.schema_agent import agent as ai_agent
from utils import _track, translate_sql, validate_table_spec, generate_sql_from_spec, _safe_json, split_sql_statements

logger = logging.getLogger(__name__)

generator_bp = Blueprint('generator', __name__)

@generator_bp.route('/create', methods=['GET', 'POST'])
def create_database():
    from services.sqlserver_service import SqlServerService
    
    # Store console logs to show in UI if something fails
    error_logs = []
    
    if request.method == 'POST':
        project_name  = request.form.get('project_name', '').strip()  # This will be the database name in SQL Server
        requirements  = request.form.get('requirements', '').strip()
        columns_spec  = request.form.get('columns_spec', '').strip()
        table_name    = request.form.get('table_name', '').strip()
        use_template  = request.form.get('use_template', '')
        db_type       = 'mssql'  # Force SQL Server

        if not project_name:
            flash('Database Name is required.', 'danger')
            return redirect(url_for('generator.create_database'))

        # Auto-sanitize project_name to a valid SQL Server identifier (remove spaces/special chars)
        sanitized_name = re.sub(r'[^A-Za-z0-9_]', '', project_name.title().replace(' ', ''))
        if sanitized_name and re.match(r'^[A-Za-z_]', sanitized_name):
            project_name = sanitized_name

        # 1. Validate Database Name (prevent SQL Injection)
        is_valid_name, err_msg = SqlServerService.validate_database_name(project_name)
        if not is_valid_name:
            flash(f"Invalid Database Name: {err_msg}", 'danger')
            return redirect(url_for('generator.create_database'))

        # 2. Check if database already exists, auto-append suffix if duplicate
        if SqlServerService.check_db_exists(project_name):
            base_name = project_name
            counter = 1
            while SqlServerService.check_db_exists(project_name):
                project_name = f"{base_name}_{counter}"
                counter += 1
            flash(f"Database '{base_name}' already exists on SQL Server. Using unique name '{project_name}'.", 'info')

        # 3. Generate tables using template, spec, or AI
        sql_code = ''
        er_mmd = ''
        entities = None
        data_dict = None
        sample_data = ''
        api_docs = None
        improvements = None
        nosql_schema = None
        ai_used = False
        
        try:
            # If using a template, load it
            if use_template:
                tmpl = IndustryTemplate.query.get(int(use_template))
                if tmpl:
                    requirements = tmpl.requirements
                    sql_code     = translate_sql(tmpl.sql_code, 'mssql')
                    er_mmd       = tmpl.er_diagram_mmd or ''
                    entities     = {'system_name': tmpl.name, 'industry': tmpl.category, 'entities': [], 'relationships': []}
                    tmpl.downloads = (tmpl.downloads or 0) + 1
                    db.session.commit()
                else:
                    flash('Template not found.', 'danger')
                    return redirect(url_for('generator.create_database'))
            elif columns_spec and table_name:
                # Direct column spec mode
                is_valid_spec, spec_err = validate_table_spec(table_name, columns_spec)
                if not is_valid_spec:
                    flash(spec_err, 'danger')
                    return redirect(url_for('generator.create_database'))
                sql_raw     = generate_sql_from_spec(table_name, columns_spec)
                sql_code    = translate_sql(sql_raw, 'mssql')
                requirements = f'Custom schema for {table_name}: {columns_spec}'
                er_mmd       = f'erDiagram\n    {table_name.upper()} {{\n        int id PK\n    }}'
                entities     = {'system_name': project_name, 'industry': 'General', 'entities': [], 'relationships': []}
            elif requirements:
                # Full AI pipeline (forces target database MSSQL)
                pipeline = ai_agent.run_full_pipeline(requirements, 'mssql')
                sql_code     = pipeline['sql_code']
                er_mmd       = pipeline['er_diagram_mmd']
                entities     = pipeline['entities']
                data_dict    = pipeline['data_dictionary']
                sample_data  = pipeline['sample_data']
                api_docs     = pipeline['api_docs']
                improvements = pipeline['improvements']
                nosql_schema = pipeline['nosql_schema']
                ai_used      = pipeline['ai_used']
            else:
                flash('Please provide requirements or column specifications.', 'danger')
                return redirect(url_for('generator.create_database'))
        except Exception as gen_err:
            flash(f"Schema generation error: {gen_err}", 'danger')
            return redirect(url_for('generator.create_database'))

        # Ensure we have SQL code
        if not sql_code:
            flash('Failed to generate table schema script.', 'danger')
            return redirect(url_for('generator.create_database'))

        # Translate SQL to Microsoft SQL Server syntax
        translated_sql = translate_sql(sql_code, 'mssql')
        statements = split_sql_statements(translated_sql)

        # 4. Execute creation in SQL Server (Database creation + Schema tables in a single operation)
        logger.info(f"Initiating SQL Server database creation and execution for: {project_name}")
        success, exec_logs = SqlServerService.create_database_and_schema(project_name, statements)
        
        if not success:
            error_logs.extend(exec_logs)
            error_msg = exec_logs[-1] if exec_logs else "Unknown database creation error."
            flash(f"SQL Server deployment failed: {error_msg}", 'danger')
            # Render create page with execution log console
            templates_featured = IndustryTemplate.query.filter_by(is_featured=True).limit(6).all()
            return render_template('create.html', 
                                   templates_featured=templates_featured,
                                   ai_enabled=cfg.AI_ENABLED,
                                   error_logs=error_logs,
                                   project_name=project_name,
                                   requirements=requirements)

        # 5. Log history to shared_databases (smartdb_system)
        industry = (entities or {}).get('industry', 'General') if isinstance(entities, dict) else 'General'
        tags_list = []
        if isinstance(entities, dict):
            for ent in entities.get('entities', []):
                tags_list.append(ent.get('name', ''))
        tags_str = ', '.join(tags_list[:5])

        new_db = SharedDatabase(
            project_name    = project_name,
            requirements    = requirements,
            sql_code        = translated_sql,
            industry        = industry,
            tags            = tags_str,
            er_diagram_mmd  = er_mmd,
            entities_json   = json.dumps(entities) if entities else None,
            data_dictionary = json.dumps(data_dict) if data_dict else None,
            sample_data     = sample_data,
            api_docs        = json.dumps(api_docs) if api_docs else None,
            improvements    = json.dumps(improvements) if improvements else None,
            nosql_schema    = json.dumps(nosql_schema) if nosql_schema else None,
            ai_generated    = ai_used,
            is_deployed     = True,
            deployed_to     = f"SQL Server ({SqlServerService.get_server()}) -> {project_name}",
            user_id         = current_user.id if current_user.is_authenticated else None
        )
        
        try:
            db.session.add(new_db)
            db.session.commit()
            _track('generated', new_db.id)
            _track('deployed', new_db.id)
            flash(f"Database '{project_name}' successfully created in SQL Server! 🚀", 'success')
            
            # Automatically connect/redirect to the newly created database (CRUD screen)
            return redirect(url_for('crud.crud_home', db_id=new_db.id))
        except Exception as db_save_err:
            db.session.rollback()
            logger.error(f"Error logging database creation history: {db_save_err}")
            flash(f"Database created in SQL Server, but failed to save history record: {db_save_err}", 'warning')
            return redirect(url_for('auth.dashboard'))
    templates_featured = IndustryTemplate.query.filter_by(is_featured=True).limit(6).all()
    return render_template('create.html', templates_featured=templates_featured,
                           ai_enabled=cfg.AI_ENABLED)



@generator_bp.route('/database/<int:db_id>')
def database_details(db_id):
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    database.views_count = (database.views_count or 0) + 1
    db.session.commit()
    _track('viewed', db_id)

    # Parse JSON fields safely
    entities_data    = _safe_json(database.entities_json)
    data_dict        = _safe_json(database.data_dictionary)
    api_docs         = _safe_json(database.api_docs)
    improvements     = _safe_json(database.improvements)
    nosql_schema     = _safe_json(database.nosql_schema)
    comments         = database.comments.order_by(Comment.created_at.desc()).all()
    versions         = database.versions.order_by(SchemaVersion.version_number.desc()).all()

    return render_template('details.html',
        database=database,
        entities_data=entities_data,
        data_dict=data_dict,
        api_docs=api_docs,
        improvements=improvements,
        nosql_schema=nosql_schema,
        comments=comments,
        versions=versions,
    )


@generator_bp.route('/database/<int:db_id>/er-diagram')
def er_diagram_view(db_id):
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    return render_template('er_diagram.html', database=database)


@generator_bp.route('/delete/<int:db_id>', methods=['POST'])
def delete_database(db_id):
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    # Check permission (owner or admin)
    if current_user.is_authenticated and (current_user.id == database.user_id or current_user.role == 'admin'):
        db.session.delete(database)
        db.session.commit()
        flash('Database entry deleted.', 'info')
    else:
        # If not signed in, check if it was anonymous (no owner)
        if database.user_id is None:
            db.session.delete(database)
            db.session.commit()
            flash('Database entry deleted.', 'info')
        else:
            flash('You do not have permission to delete this database.', 'danger')
    return redirect(url_for('marketplace.store'))


@generator_bp.route('/download-sql/<int:db_id>')
def download_sql_file(db_id):
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    safe     = re.sub(r'[^a-zA-Z0-9_]', '_', database.project_name.lower())
    log      = DownloadLog(schema_id=db_id, ip_address=request.remote_addr, file_type='sql',
                           user_id=current_user.id if current_user.is_authenticated else None)
    database.downloads_count = (database.downloads_count or 0) + 1
    db.session.add(log)
    db.session.commit()
    _track('downloaded', db_id)
    return Response(database.sql_code, mimetype='text/plain',
                    headers={'Content-Disposition': f'attachment; filename={safe}.sql'})


@generator_bp.route('/execute/<int:db_id>', methods=['GET', 'POST'])
def execute_schema(db_id):
    database     = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    logs         = []
    success      = False
    download_url = None
    db_type      = 'sqlite'
    host='localhost'; port=''; username=''; password=''; db_name=''

    if request.method == 'POST':
        db_type  = request.form.get('db_type', 'sqlite')
        host     = request.form.get('host', 'localhost')
        port     = request.form.get('port', '')
        username = request.form.get('username', '')
        password = request.form.get('password', '')
        db_name  = request.form.get('db_name', '').strip()

        try:
            logs.append(f"Translating SQL schema dialect for target DBMS: {db_type.upper()}")
            translated_sql = translate_sql(database.sql_code, db_type)

            if db_type == 'sqlite':
                safe_proj  = re.sub(r'[^a-zA-Z0-9_]', '_', database.project_name.lower())
                db_fname   = f"{safe_proj}_{db_id}.db"
                db_fpath   = os.path.join(cfg.SQLITE_DIR, db_fname)
                if os.path.exists(db_fpath):
                     os.remove(db_fpath)
                conn   = sqlite3.connect(db_fpath)
                cursor = conn.cursor()
                logs.append(f"Initialising SQLite file: {db_fname}")
                stmts = split_sql_statements(translated_sql)
                for idx, stmt in enumerate(stmts, 1):
                    logs.append(f"Executing statement {idx}: {stmt[:70]}...")
                    cursor.execute(stmt)
                conn.commit(); conn.close()
                success      = True
                download_url = url_for('static', filename=f'generated_databases/{db_fname}')
                logs.append("All statements executed. Database file ready for download.")
                database.downloads_count = (database.downloads_count or 0) + 1
                database.is_deployed = True
                database.deployed_to = "SQLite Local File"
                db.session.commit()
                _track('deployed', db_id)
            else:
                if not db_name:
                    raise ValueError("Database Name is required for server deployments.")
                from sqlalchemy import create_engine, text
                if not port:
                    port = {'mysql':'3306','postgresql':'5432','mssql':'1433'}.get(db_type,'5432')
                if db_type == 'mysql':
                    import pymysql
                    conn_url  = f"mysql+pymysql://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/{db_name}"
                    admin_url = f"mysql+pymysql://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/"
                elif db_type == 'postgresql':
                    import psycopg2
                    conn_url  = f"postgresql+psycopg2://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/{db_name}"
                    admin_url = f"postgresql+psycopg2://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/postgres"
                elif db_type == 'mssql':
                    import pyodbc
                    drv       = 'ODBC Driver 17 for SQL Server'
                    conn_url  = f"mssql+pyodbc://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/{db_name}?driver={quote_plus(drv)}&Encrypt=no"
                    admin_url = f"mssql+pyodbc://{quote_plus(username)}:{quote_plus(password)}@{host}:{port}/master?driver={quote_plus(drv)}&Encrypt=no"
                logs.append(f"Connecting to {db_type.upper()} at {host}:{port}...")
                try:
                    ae = create_engine(admin_url)
                    with ae.execution_options(isolation_level='AUTOCOMMIT').connect() as c:
                        if db_type == 'mysql':
                            c.execute(text(f"CREATE DATABASE IF NOT EXISTS {db_name}"))
                        elif db_type == 'postgresql':
                            res = c.execute(text(f"SELECT 1 FROM pg_database WHERE datname='{db_name}'"))
                            if not res.scalar():
                                c.execute(text(f"CREATE DATABASE {db_name}"))
                        logs.append(f"Database '{db_name}' ready.")
                    ae.dispose()
                except Exception as e:
                    logs.append(f"Auto-create note: {e}")
                engine = create_engine(conn_url)
                with engine.execution_options(isolation_level='AUTOCOMMIT').connect() as c:
                    stmts = split_sql_statements(translated_sql)
                    for idx, stmt in enumerate(stmts, 1):
                        logs.append(f"Executing {idx}: {stmt[:70]}...")
                        c.execute(text(stmt))
                engine.dispose()
                logs.append("All statements executed successfully!")
                success = True
                database.is_deployed = True
                database.deployed_to = f"{db_type.upper()} ({host}:{port}) -> {db_name}"
                db.session.commit()
                _track('deployed', db_id)
        except Exception as e:
            logs.append(f"ERROR: {e}")
            success = False

    return render_template('deploy.html', database=database, logs=logs, success=success,
                           download_url=download_url, db_type=db_type,
                           host=host, port=port, username=username, password=password, db_name=db_name)


@generator_bp.route('/push-to-ssms/<int:db_id>', methods=['GET', 'POST'])
def push_to_ssms(db_id):
    from services.sqlserver_service import SqlServerService
    database    = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort
        abort(404)
    logs        = []
    success     = False
    db_name     = ''
    server_name = SqlServerService.get_server()
    auth_mode   = 'windows'
    sql_user     = ''
    sql_password = ''

    if request.method == 'POST':
        db_name      = request.form.get('db_name', '').strip()
        server_name  = request.form.get('server_name', SqlServerService.get_server()).strip()
        auth_mode    = request.form.get('auth_mode', 'windows')
        sql_user     = request.form.get('sql_user', '').strip()
        sql_password = request.form.get('sql_password', '')

        if not db_name:
            logs.append("ERROR: Database name is required.")
        else:
            try:
                import pyodbc
                available = [d for d in pyodbc.drivers() if 'SQL Server' in d]
                if not available:
                    raise RuntimeError("No SQL Server ODBC driver found. Install 'ODBC Driver 17 for SQL Server'.")
                chosen = next((d for d in ['ODBC Driver 18 for SQL Server','ODBC Driver 17 for SQL Server','SQL Server'] if d in available), available[0])
                logs.append(f"Using driver: {chosen}")

                def build_conn_str(srv, db='master'):
                    base = f"DRIVER={{{chosen}}};SERVER={srv};DATABASE={db};Encrypt=no;TrustServerCertificate=yes;"
                    return base + ("Trusted_Connection=yes;" if auth_mode=='windows' else f"UID={sql_user};PWD={sql_password};")

                master = None
                try:
                    master = pyodbc.connect(build_conn_str(server_name, 'master'), autocommit=True)
                except pyodbc.Error as conn_err:
                    # If connecting to plain 'localhost' fails, auto-try system configured instance (e.g. localhost\SQLEXPRESS04)
                    sys_server = SqlServerService.get_server()
                    if server_name != sys_server:
                        logs.append(f"Could not connect to '{server_name}'. Retrying with detected instance '{sys_server}'...")
                        try:
                            master = pyodbc.connect(build_conn_str(sys_server, 'master'), autocommit=True)
                            server_name = sys_server
                            logs.append(f"Connected successfully to '{sys_server}'.")
                        except pyodbc.Error:
                            raise conn_err
                    else:
                        raise conn_err

                cur = master.cursor()
                cur.execute("SELECT COUNT(*) FROM sys.databases WHERE name=?", db_name)
                if cur.fetchone()[0]:
                    logs.append(f"Database '{db_name}' exists — using it.")
                else:
                    cur.execute(f"CREATE DATABASE [{db_name}]")
                    logs.append(f"Created database '{db_name}'.")
                master.close()

                translated = translate_sql(database.sql_code, 'mssql')
                target = pyodbc.connect(build_conn_str(server_name, db_name), autocommit=True)
                tcur   = target.cursor()
                stmts  = split_sql_statements(translated)
                for idx, stmt in enumerate(stmts, 1):
                    logs.append(f"Executing {idx}: {stmt[:70]}...")
                    tcur.execute(stmt)
                target.commit(); target.close()
                logs.append(f"Done! '{db_name}' is live on SQL Server. Open SSMS → Databases → Refresh.")
                success = True
                database.is_deployed = True
                database.deployed_to = f"SQL Server ({server_name}) -> {db_name}"
                db.session.commit()
                _track('deployed', db_id)
            except ImportError:
                logs.append("ERROR: 'pyodbc' not installed. Run: pip install pyodbc")
            except Exception as e:
                logs.append(f"ERROR: {e}")

    return render_template('ssms.html', database=database, logs=logs, success=success,
                           db_name=db_name, server_name=server_name, auth_mode=auth_mode, sql_user=sql_user)


@generator_bp.route('/detect-sqlserver')
def detect_sqlserver():
    try:
        import pyodbc
        drivers = [d for d in pyodbc.drivers() if 'SQL Server' in d]
        return jsonify({'drivers': drivers, 'available': len(drivers) > 0})
    except ImportError:
        return jsonify({'drivers': [], 'available': False, 'error': 'pyodbc not installed'})


# ── Export Routes ─────────────────────────────────────────────────────────────

@generator_bp.route('/export-csv/<int:db_id>')
def export_csv(db_id):
    """Export dynamic table data from SQL Server as CSV."""
    from services.sqlserver_service import SqlServerService
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort; abort(404)

    db_name = database.project_name
    if not SqlServerService.check_db_exists(db_name):
        flash(f"SQL Server database '{db_name}' not found.", 'warning')
        return redirect(url_for('generator.database_details', db_id=db_id))

    table_name = request.args.get('table', '')
    try:
        conn = SqlServerService.get_connection(db_name)
        tables = SqlServerService.get_tables(db_name)
        
        if not table_name:
            table_name = tables[0] if tables else None
            
        if not table_name:
            conn.close()
            flash('No user tables found in database.', 'danger')
            return redirect(url_for('generator.database_details', db_id=db_id))
            
        if table_name not in tables:
            conn.close()
            abort(404)
            
        cursor = conn.cursor()
        cursor.execute(f"SELECT * FROM [{table_name}]")
        col_names = [col[0] for col in cursor.description]
        rows = [dict(zip(col_names, row)) for row in cursor.fetchall()]
        cursor.close()
        conn.close()
    except Exception as e:
        flash(f"Failed to read data from SQL Server: {e}", 'danger')
        return redirect(url_for('generator.database_details', db_id=db_id))

    output = io.StringIO()
    if rows:
        writer = csv.DictWriter(output, fieldnames=rows[0].keys())
        writer.writeheader()
        for row in rows:
            writer.writerow(row)

    _track('downloaded', db_id)
    safe_table = re.sub(r'[^a-zA-Z0-9_]', '_', table_name.lower())
    filename = f"{db_name}_{safe_table}.csv"
    return Response(
        output.getvalue(),
        mimetype='text/csv',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )


@generator_bp.route('/export-json/<int:db_id>')
def export_json(db_id):
    """Export all tables from SQL Server as a single JSON file."""
    from services.sqlserver_service import SqlServerService
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort; abort(404)

    db_name = database.project_name
    if not SqlServerService.check_db_exists(db_name):
        flash(f"SQL Server database '{db_name}' not found.", 'warning')
        return redirect(url_for('generator.database_details', db_id=db_id))

    try:
        conn = SqlServerService.get_connection(db_name)
        tables = SqlServerService.get_tables(db_name)
        
        all_data = {}
        cursor = conn.cursor()
        for tbl in tables:
            cursor.execute(f"SELECT * FROM [{tbl}]")
            col_names = [col[0] for col in cursor.description]
            all_data[tbl] = [dict(zip(col_names, row)) for row in cursor.fetchall()]
        cursor.close()
        conn.close()
    except Exception as e:
        flash(f"Failed to extract JSON data from SQL Server: {e}", 'danger')
        return redirect(url_for('generator.database_details', db_id=db_id))

    _track('downloaded', db_id)
    filename = f"{db_name}_export.json"
    return Response(
        json.dumps(all_data, indent=2, default=str),
        mimetype='application/json',
        headers={'Content-Disposition': f'attachment; filename={filename}'}
    )


@generator_bp.route('/export-excel/<int:db_id>')
def export_excel(db_id):
    """Export all tables and records from SQL Server to a formatted multi-sheet Excel file."""
    from services.sqlserver_service import SqlServerService
    database = db.session.get(SharedDatabase, db_id)
    if database is None:
        from flask import abort; abort(404)

    db_name = database.project_name
    if not SqlServerService.check_db_exists(db_name):
        flash(f"SQL Server database '{db_name}' not found.", 'warning')
        return redirect(url_for('generator.database_details', db_id=db_id))

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except ImportError:
        flash('openpyxl package is not installed. Run: pip install openpyxl', 'danger')
        return redirect(url_for('generator.database_details', db_id=db_id))

    try:
        conn = SqlServerService.get_connection(db_name)
        tables = SqlServerService.get_tables(db_name)
        
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        # Premium Style definitions
        header_fill = PatternFill('solid', fgColor='003366')
        header_font = Font(bold=True, color='FFFFFF', name='Calibri', size=11)
        body_font   = Font(name='Calibri', size=10)
        thin_border = Border(
            left=Side(style='thin', color='DDDDDD'), right=Side(style='thin', color='DDDDDD'),
            top=Side(style='thin', color='DDDDDD'), bottom=Side(style='thin', color='DDDDDD')
        )

        cursor = conn.cursor()
        for tbl in tables:
            ws = wb.create_sheet(title=tbl[:31])  # sheet name max 31 chars
            cursor.execute(f"SELECT * FROM [{tbl}]")
            col_names = [col[0] for col in cursor.description]
            rows = cursor.fetchall()
            
            # Write Headers
            ws.append(col_names)
            for cell in ws[1]:
                cell.font      = header_font
                cell.fill      = header_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border    = thin_border

            # Write Rows
            for row in rows:
                ws.append(list(row))
                for cell in ws[ws.max_row]:
                    cell.font   = body_font
                    cell.border = thin_border

            # Auto-fit column widths
            for idx, col in enumerate(ws.columns, 1):
                max_len = max((len(str(cell.value or '')) for cell in col), default=10)
                ws.column_dimensions[get_column_letter(idx)].width = min(max_len + 4, 40)
                
        cursor.close()
        conn.close()
    except Exception as e:
        flash(f"Failed to generate Excel sheet from SQL Server: {e}", 'danger')
        return redirect(url_for('generator.database_details', db_id=db_id))

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    _track('downloaded', db_id)

    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=f"{db_name}_data.xlsx"
    )


# ── Clone Route ───────────────────────────────────────────────────────────────

@generator_bp.route('/clone/<int:db_id>', methods=['POST'])
def clone_database(db_id):
    """Clone (duplicate) an existing schema as a new project."""
    source = db.session.get(SharedDatabase, db_id)
    if source is None:
        from flask import abort; abort(404)

    new_name = request.form.get('clone_name', f'{source.project_name} (Clone)').strip()

    clone = SharedDatabase(
        project_name    = new_name,
        requirements    = source.requirements,
        sql_code        = source.sql_code,
        industry        = source.industry,
        tags            = source.tags,
        er_diagram_mmd  = source.er_diagram_mmd,
        entities_json   = source.entities_json,
        data_dictionary = source.data_dictionary,
        sample_data     = source.sample_data,
        api_docs        = source.api_docs,
        improvements    = source.improvements,
        nosql_schema    = source.nosql_schema,
        ai_generated    = source.ai_generated,
        user_id         = current_user.id if current_user.is_authenticated else None,
        is_deployed     = False,
        deployed_to     = None,
    )
    db.session.add(clone)
    db.session.commit()
    _track('generated', clone.id)
    flash(f'Schema cloned as "{new_name}"! 🎉 You can now modify and deploy it.', 'success')
    return redirect(url_for('generator.database_details', db_id=clone.id))
